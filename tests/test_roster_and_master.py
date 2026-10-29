"""참여자 명단 업로드 · 개인 참여 코드 · 마스터 대리(고객사로 들어가기) · 고객사별 현황 · 접두어 규칙.
실행: pytest tests/test_roster_and_master.py
"""
import io
import uuid

from openpyxl import Workbook, load_workbook

from diag_project.services import org_codes
from tests.test_org_units import _client_admin_headers, _new_company

BASE = "/api/v1/admin"


def _xlsx(rows, header=("이름", "이메일", "본부", "실", "팀")):
    wb = Workbook()
    ws = wb.active
    ws.append(["참여자 명단"])          # 제목 줄이 있어도 머리글을 찾는다
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, headers, cid, content, dry_run):
    return client.post(f"{BASE}/companies/{cid}/roster", params={"dry_run": str(dry_run).lower()},
                       files={"file": ("roster.xlsx", content,
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
                       headers=headers)


def test_prefix_from_company_code_or_name():
    assert org_codes.prefix_candidates(["HICARE-01", "하이케어솔루션"]) == ["H", "HI", "HIC"]
    assert org_codes.auto_prefix([], hints=["HICARE-01"]) == "H"
    assert org_codes.auto_prefix(["H"], hints=["HICARE-01"]) == "A"        # H 와 HI 는 앞부분이 겹쳐서 건너뜀 → 순서대로
    assert org_codes.auto_prefix(["X"], hints=["07960", "커넥트"]) == "A"   # 영문이 없으면 A, B, …


def test_roster_upload_sets_org_and_personal_codes(client, auth_headers):
    code = "RST" + uuid.uuid4().hex[:5].upper()
    cid = _new_company(client, auth_headers, code)
    other = _new_company(client, auth_headers, "OTH-" + uuid.uuid4().hex[:6])
    tag = uuid.uuid4().hex[:6]
    rows = [
        ("홍길동", f"gd-{tag}@x.com", "영업본부", "", "영업1팀"),
        ("김하늘", f"hn-{tag}@x.com", "영업본부", "마케팅실", "브랜드팀"),
        ("이도현", f"dh-{tag}@x.com", "생산본부", "", ""),
        ("박미분류", f"pm-{tag}@x.com", "", "", ""),
    ]
    hr = _client_admin_headers(client, cid)

    # 미리보기: 저장 안 함
    r = _upload(client, hr, cid, _xlsx(rows), dry_run=True)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["saved"] is False and d["new"] == 4 and not d["errors"]
    assert client.get(f"{BASE}/companies/{cid}/org-units", headers=hr).json()["divisions"] == []

    # 확정: 조직 자동 생성 + 개인 코드
    d = _upload(client, hr, cid, _xlsx(rows), dry_run=False).json()
    assert d["saved"] is True and d["new"] == 4
    tree = client.get(f"{BASE}/companies/{cid}/org-units", headers=hr).json()
    prefix = tree["company"]["org_prefix"]
    assert prefix and prefix[0] == "R"                     # 회사 코드(RST…)의 영문에서 접두어
    names = {x["name"] for x in tree["divisions"]}
    assert names == {"영업본부", "생산본부"}

    lst = client.get(f"{BASE}/participants", params={"page_size": 50}, headers=hr).json()["items"]
    by_mail = {i["email"]: i for i in lst}
    assert by_mail[f"hn-{tag}@x.com"]["office_name"] == "마케팅실"
    assert by_mail[f"pm-{tag}@x.com"]["org_code"] is None
    assert all(i["progress_status"] == "not_started" for i in lst)

    # 개인 코드 엑셀
    x = client.get(f"{BASE}/companies/{cid}/roster/export", headers=hr)
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content))["개인 참여 코드"]
    table = [r for r in ws.iter_rows(min_row=7, values_only=True) if r[0]]
    codes = {r[1]: r[5] for r in table}
    gd_code = codes[f"gd-{tag}@x.com"]
    assert gd_code.startswith(prefix) and "-" in gd_code and len(gd_code.split("-")[1]) == 4

    # 개인 코드로 로그인: 이메일이 맞으면 그 사람, 틀리면 거부
    ok = client.post("/api/v1/participants/token", json={"email": f"gd-{tag}@x.com", "name": "홍길동", "group_code": gd_code.lower(), "password": "x"})
    assert ok.status_code == 200 and ok.json()["participant_id"] == by_mail[f"gd-{tag}@x.com"]["id"]
    bad = client.post("/api/v1/participants/token", json={"email": f"someone-{tag}@x.com", "name": "x", "group_code": gd_code, "password": "x"})
    assert bad.status_code == 403

    # 다시 올리면 갱신(소속이 같으면 코드 유지), 오류 행이 있으면 아무것도 저장 안 함
    again = _upload(client, hr, cid, _xlsx(rows[:1]), dry_run=False).json()
    assert again["updated"] == 1 and again["new"] == 0
    x2 = load_workbook(io.BytesIO(client.get(f"{BASE}/companies/{cid}/roster/export", headers=hr).content))["개인 참여 코드"]
    assert {r[1]: r[5] for r in x2.iter_rows(min_row=7, values_only=True) if r[0]}[f"gd-{tag}@x.com"] == gd_code

    bad_rows = rows[:1] + [("", "nope", "영업본부", "", ""), ("중복", f"gd-{tag}@x.com", "", "", "")]
    e = _upload(client, hr, cid, _xlsx(bad_rows), dry_run=False).json()
    assert e["saved"] is False and {x["row"] for x in e["errors"]} == {4, 5}

    # 다른 고객사에 이미 있는 이메일 → 오류
    e2 = _upload(client, auth_headers, other, _xlsx(rows[:1]), dry_run=True).json()
    assert e2["errors"] and "다른 고객사" in e2["errors"][0]["reason"]

    # 머리글 없는 파일 → 400, 다른 회사 HR → 403, 양식 다운로드
    assert _upload(client, hr, cid, _xlsx(rows, header=("a", "b", "c", "d", "e")), dry_run=True).status_code == 400
    assert _upload(client, hr, other, _xlsx(rows), dry_run=True).status_code == 403
    assert client.get(f"{BASE}/roster/template", headers=hr).status_code == 200


def test_master_acting_and_overview(client, auth_headers):
    a = _new_company(client, auth_headers, "ACT-" + uuid.uuid4().hex[:6])
    b = _new_company(client, auth_headers, "ACT-" + uuid.uuid4().hex[:6])
    tag = uuid.uuid4().hex[:6]
    _upload(client, auth_headers, a, _xlsx([("에이", f"a-{tag}@x.com", "", "", "")]), dry_run=False)
    _upload(client, auth_headers, b, _xlsx([("비", f"b-{tag}@x.com", "", "", "")]), dry_run=False)

    # 고객사로 들어가기: me 는 그 고객사, 목록·통계도 그 고객사만
    act = {**auth_headers, "X-Act-As-Company": a}
    me = client.get(f"{BASE}/auth/me", headers=act).json()
    assert me["acting"] is True and me["company_id"] == a and me["role"] == "super_admin"
    emails = {i["email"] for i in client.get(f"{BASE}/participants", params={"page_size": 200}, headers=act).json()["items"]}
    assert f"a-{tag}@x.com" in emails and f"b-{tag}@x.com" not in emails
    assert client.get(f"{BASE}/stats/overview", headers=act).json()["scope"] == "company"

    # 없는 고객사 → 404, HR 이 헤더를 보내도 무시(권한 상승 불가)
    assert client.get(f"{BASE}/auth/me", headers={**auth_headers, "X-Act-As-Company": str(uuid.uuid4())}).status_code == 404
    hr_a = _client_admin_headers(client, a)
    me_hr = client.get(f"{BASE}/auth/me", headers={**hr_a, "X-Act-As-Company": b}).json()
    assert me_hr["company_id"] == a and me_hr["acting"] is False

    # 고객사별 현황(마스터 전용)
    ov = client.get(f"{BASE}/companies/overview", headers=auth_headers)
    assert ov.status_code == 200
    row = next(r for r in ov.json()["companies"] if r["id"] == a)
    assert row["registered"] == 1 and row["roster"] == 1 and row["not_started"] == 1 and row["completion_rate"] == 0
    assert client.get(f"{BASE}/companies/overview", headers=hr_a).status_code == 403
