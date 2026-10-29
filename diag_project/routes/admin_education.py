# diag_project/routes/admin_education.py
#
# 고객사 HR 콘솔의 '결과 → 교육' 데이터.
#   - GET /admin/education/summary : 강화·개선 필요 역량 분포, 교육 체계도(역량 × 트랙 과정별 추천 인원), 본부 × 역량 평균
#   - GET /admin/education/people  : 개인별 강화·개선 필요 역량과 추천 과정 명단
# 모든 값은 저장된 리포트(scores.radar_chart · details · course_recommendation)에서만 계산한다(지어내지 않음).
# 회사 격리는 AdminContext.scope_query 로 강제한다.

from collections import defaultdict
from typing import Any, Dict, List, Optional, Tuple
from uuid import UUID

from fastapi import APIRouter, Depends, Query
from sqlalchemy import or_
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select

from diag_project.data.competencies import COMPETENCY_FRAMEWORK
from diag_project.database import get_db
from diag_project.models.diagnosis_report import DiagnosisReport
from diag_project.models.org_unit import OrgUnit
from diag_project.models.participant import Participant
from diag_project.services.auth import AdminContext, get_current_admin
from diag_project.services.org_units import descendant_ids, unit_labels

router = APIRouter(prefix="/admin/education", tags=["Admin · 교육"])

_ORDER = list(COMPETENCY_FRAMEWORK.keys())
_NAMES = {k: v.get("name", k) for k, v in COMPETENCY_FRAMEWORK.items()}
# 사내 차수 개설을 권하는 추천 인원 기준(교육 체계도 화면의 강조 기준과 같다)
COHORT_MIN = 15


def _num(v: Any) -> Optional[float]:
    try:
        return float(v) if v is not None else None
    except (TypeError, ValueError):
        return None


def measured_scores(scores: Dict[str, Any]) -> Dict[str, float]:
    """대역량 키 → 측정 점수. details 에 measured_count 가 있으면 그 기준(0이면 미측정), 없으면(구 리포트) radar 값."""
    radar = scores.get("radar_chart") or {}
    details = scores.get("details") or {}
    out: Dict[str, float] = {}
    for k in _ORDER:
        d = details.get(k) or {}
        if "measured_count" in d and int(d.get("measured_count") or 0) <= 0:
            continue
        v = _num(d.get("score")) if "measured_count" in d else _num(radar.get(k, radar.get(_NAMES[k])))
        if v is not None:
            out[k] = v
    return out


def strength_and_growth(scores: Dict[str, Any]) -> Tuple[Optional[str], Optional[str]]:
    """강화 역량(최고)·개선 필요 역량(최저). 점수가 같으면 하위역량 평균으로 가른다. 측정 2개 미만이면 (None, None).

    리포트 화면(ReportNextStep.tsx)의 rank() 와 같은 규칙.
    """
    ms = measured_scores(scores)
    if len(ms) < 2:
        return None, None
    details = scores.get("details") or {}

    def sub_avg(k: str) -> float:
        subs = [x for x in ((details.get(k) or {}).get("sub_scores") or {}).values() if isinstance(x, (int, float))]
        return sum(subs) / len(subs) if subs else 0.0

    ranked = sorted(ms.keys(), key=lambda k: (ms[k], sub_avg(k)), reverse=True)
    return ranked[0], ranked[-1]


def courses_of(scores: Dict[str, Any]) -> List[Dict[str, Any]]:
    """리포트의 추천 과정 목록(개선 과정 + 강점 활용 과정). 없으면 빈 목록."""
    cr = scores.get("course_recommendation") or {}
    out = []
    for c in (cr.get("growth") or []):
        if isinstance(c, dict) and c.get("course"):
            out.append({**c, "kind": "growth"})
    s = cr.get("strength")
    if isinstance(s, dict) and s.get("course"):
        out.append({**s, "kind": "strength"})
    return out


async def _scoped_participants(db: AsyncSession, ctx: AdminContext, company_id: Optional[UUID],
                               org_unit_id: Optional[UUID]):
    q = ctx.scope_query(select(Participant), Participant.company_id)
    if company_id and ctx.is_super_admin:
        q = q.where(Participant.company_id == company_id)
    if org_unit_id:
        q = q.where(Participant.org_unit_id.in_(await descendant_ids(db, org_unit_id)))
    return q


async def _latest_reports(db: AsyncSession, participant_ids: List[UUID]) -> Dict[UUID, DiagnosisReport]:
    if not participant_ids:
        return {}
    rows = (await db.execute(
        select(DiagnosisReport).where(DiagnosisReport.user_id.in_(participant_ids))
        .order_by(DiagnosisReport.created_at.desc())
    )).scalars().all()
    out: Dict[UUID, DiagnosisReport] = {}
    for r in rows:
        out.setdefault(r.user_id, r)
    return out


@router.get("/summary")
async def education_summary(
    company_id: Optional[UUID] = Query(None, description="마스터 전용 고객사 필터"),
    org_unit_id: Optional[UUID] = Query(None, description="조직 단위 필터(본부면 소속 팀 포함)"),
    ctx: AdminContext = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db),
):
    participants = (await db.execute(
        await _scoped_participants(db, ctx, company_id, org_unit_id)
    )).scalars().all()
    reports = await _latest_reports(db, [p.id for p in participants])

    strength_n: Dict[str, int] = defaultdict(int)
    growth_n: Dict[str, int] = defaultdict(int)
    courses: Dict[Tuple[str, str, str], Dict[str, Any]] = {}
    # 본부 × 역량 평균용: 본부 id → 역량 → 점수 목록
    by_div: Dict[Optional[str], Dict[str, List[float]]] = defaultdict(lambda: defaultdict(list))
    labels = await unit_labels(db, [p.org_unit_id for p in participants])

    for p in participants:
        rep = reports.get(p.id)
        if rep is None:
            continue
        scores = rep.scores or {}
        s, g = strength_and_growth(scores)
        if s:
            strength_n[s] += 1
        if g:
            growth_n[g] += 1
        for c in courses_of(scores):
            key = (c.get("category") or "", c.get("track") or "", c.get("course") or "")
            row = courses.setdefault(key, {
                "category": key[0], "track": key[1], "course": key[2],
                "track_format": c.get("track_format"), "vod_url": c.get("vod_url"),
                "growth": 0, "strength": 0,
            })
            row[c["kind"]] += 1
        div = labels.get(p.org_unit_id, {}).get("division_id") if p.org_unit_id else None
        for k, v in measured_scores(scores).items():
            by_div[div][k].append(v)

    completed = len(reports)
    course_rows = sorted(courses.values(), key=lambda r: (-(r["growth"] + r["strength"]), r["category"], r["track"]))
    for r in course_rows:
        r["total"] = r["growth"] + r["strength"]
        r["cohort"] = r["total"] >= COHORT_MIN

    divisions = []
    if by_div:
        div_units = {}
        div_ids = [UUID(d) for d in by_div if d]
        if div_ids:
            div_units = {str(u.id): u for u in (await db.execute(select(OrgUnit).where(OrgUnit.id.in_(div_ids)))).scalars().all()}
        for d, comp in by_div.items():
            u = div_units.get(d) if d else None
            divisions.append({
                "id": d, "name": u.name if u else "미분류", "code": u.code if u else None,
                "n": max((len(v) for v in comp.values()), default=0),
                "averages": {k: (round(sum(v) / len(v), 1) if v else None) for k, v in ((k, comp.get(k, [])) for k in _ORDER)},
            })
        divisions.sort(key=lambda x: (x["id"] is None, x["code"] or ""))

    return {
        "competencies": [{"key": k, "name": _NAMES[k]} for k in _ORDER],
        "participants": len(participants),
        "completed": completed,
        "strength": [{"key": k, "name": _NAMES[k], "count": strength_n.get(k, 0)} for k in _ORDER],
        "growth": [{"key": k, "name": _NAMES[k], "count": growth_n.get(k, 0)} for k in _ORDER],
        "courses": course_rows,
        "cohort_min": COHORT_MIN,
        "divisions": divisions,
    }


@router.get("/people")
async def education_people(
    company_id: Optional[UUID] = Query(None, description="마스터 전용 고객사 필터"),
    org_unit_id: Optional[UUID] = Query(None),
    search: Optional[str] = Query(None),
    status: Optional[str] = Query(None, description="completed | not_completed"),
    course: Optional[str] = Query(None, description="추천 과정명으로 거르기"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=200),
    ctx: AdminContext = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db),
):
    q = await _scoped_participants(db, ctx, company_id, org_unit_id)
    if search:
        kw = f"%{search.strip()}%"
        q = q.where(or_(Participant.name.ilike(kw), Participant.email.ilike(kw)))
    participants = (await db.execute(q.order_by(Participant.name))).scalars().all()
    reports = await _latest_reports(db, [p.id for p in participants])
    labels = await unit_labels(db, [p.org_unit_id for p in participants])

    rows = []
    for p in participants:
        rep = reports.get(p.id)
        scores = (rep.scores or {}) if rep else {}
        s, g = strength_and_growth(scores) if rep else (None, None)
        recs = [{"track": c.get("track"), "course": c.get("course"), "category": c.get("category"),
                 "kind": c["kind"], "vod_url": c.get("vod_url")} for c in courses_of(scores)] if rep else []
        if status == "completed" and rep is None:
            continue
        if status == "not_completed" and rep is not None:
            continue
        if course and not any(r["course"] == course for r in recs):
            continue
        org = labels.get(p.org_unit_id) if p.org_unit_id else None
        rows.append({
            "id": str(p.id), "name": p.name or "-", "email": p.email,
            "org_code": org["code"] if org else None, "org_name": org["name"] if org else None,
            "division_name": org["division_name"] if org else None,
            "completed": rep is not None, "report_id": str(rep.id) if rep else None,
            "total_score": rep.total_score if rep else None,
            "strength": _NAMES.get(s) if s else None, "growth": _NAMES.get(g) if g else None,
            "courses": recs,
        })

    total = len(rows)
    start = (page - 1) * page_size
    return {"items": rows[start:start + page_size], "total": total, "page": page, "page_size": page_size,
            "total_pages": max(1, (total + page_size - 1) // page_size)}


@router.get("/catalog")
async def education_catalog(
    company_id: Optional[UUID] = Query(None, description="마스터 전용 고객사 필터"),
    org_unit_id: Optional[UUID] = Query(None),
    ctx: AdminContext = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db),
):
    """교육 체계도·과정 안내서: 전체 과정 카탈로그(역량 → 하위역량 × 트랙 A~D) + 범위 안 추천 인원.

    과정명·교육목표(수료 후 행동)·학습 방식은 data/course_matrix.py 그대로. 추천 인원은 저장된 리포트의
    course_recommendation 을 (역량, 하위역량, 트랙)으로 맞춰 센다.
    """
    from diag_project.data.course_matrix import COURSE_CATALOG, TRACK_META, VOD_BASE

    participants = (await db.execute(
        await _scoped_participants(db, ctx, company_id, org_unit_id)
    )).scalars().all()
    reports = await _latest_reports(db, [p.id for p in participants])
    counts: Dict[Tuple[str, str, str], Dict[str, int]] = defaultdict(lambda: {"growth": 0, "strength": 0})
    for rep in reports.values():
        for c in courses_of(rep.scores or {}):
            counts[(c.get("category") or "", c.get("sub_competency") or "", c.get("track") or "")][c["kind"]] += 1

    comps = []
    for ck in _ORDER:
        subs = []
        for (k, sub_name), entry in COURSE_CATALOG.items():
            if k != ck:
                continue
            courses = {}
            for t in ("A", "B", "C", "D"):
                c = entry.get(t)
                if not c:
                    continue
                n = counts.get((_NAMES[ck], sub_name, t), {"growth": 0, "strength": 0})
                total = n["growth"] + n["strength"]
                courses[t] = {
                    "course": c.get("course"), "goal": c.get("subtitle"),
                    # VOD 주소가 아직 자리표시(example.com)면 내보내지 않는다(깨진 링크 방지)
                    "vod_url": (f"{VOD_BASE}/{entry.get('slug')}-{t.lower()}"
                                if entry.get("slug") and "example.com" not in VOD_BASE else None),
                    "growth": n["growth"], "strength": n["strength"], "total": total, "cohort": total >= COHORT_MIN,
                }
            subs.append({"name": sub_name, "courses": courses})
        comps.append({"key": ck, "name": _NAMES[ck], "subs": subs})

    tracks = [{
        "track": t, "from_level": m.get("diagnosed_level"), "to_level": m.get("target_level"),
        "format": m.get("format"), "format_desc": m.get("format_desc"),
    } for t, m in TRACK_META.items()]
    return {"tracks": tracks, "competencies": comps, "completed": len(reports), "cohort_min": COHORT_MIN}


@router.get("/curriculum")
async def education_curriculum(
    company_id: Optional[UUID] = Query(None, description="마스터 전용 고객사"),
    org_unit_id: Optional[UUID] = Query(None, description="본부·실·팀(고르면 그 조직 결과로 구성, 회사 전체와 비교)"),
    ctx: AdminContext = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db),
):
    """맞춤 교육체계도: 고른 범위의 진단 결과로 Pool 에서 필수·선택·강점 확산 과정을 골라 구성(services/curriculum.py)."""
    from diag_project.services.curriculum import build_curriculum

    participants = (await db.execute(
        await _scoped_participants(db, ctx, company_id, org_unit_id)
    )).scalars().all()
    reports = await _latest_reports(db, [p.id for p in participants])
    details = [((r.scores or {}).get("details")) or {} for r in reports.values()]

    baseline = None
    if org_unit_id:   # 조직을 골랐으면 같은 회사 전체를 비교 기준으로
        all_ps = (await db.execute(await _scoped_participants(db, ctx, company_id, None))).scalars().all()
        all_reports = await _latest_reports(db, [p.id for p in all_ps])
        baseline = [((r.scores or {}).get("details")) or {} for r in all_reports.values()]
    return build_curriculum(details, baseline)


@router.get("/curriculum/export")
async def export_curriculum(
    company_id: Optional[UUID] = Query(None, description="마스터 전용 고객사"),
    org_unit_id: Optional[UUID] = Query(None),
    ctx: AdminContext = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db),
):
    """맞춤 교육체계도 엑셀: 화면 포스터와 같은 구성(남색 머리 띠·회색 트랙 머리·역량 세로 병합·필수/선택/강점 색)."""
    import io
    from datetime import datetime
    from urllib.parse import quote

    from fastapi import HTTPException
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    from diag_project.models.company import Company
    from diag_project.services.curriculum import build_curriculum
    from diag_project.services.org_units import unit_labels

    cur = await education_curriculum(company_id=company_id, org_unit_id=org_unit_id, ctx=ctx, db=db)
    if cur["status"] != "ok":
        raise HTTPException(status_code=400, detail=f"진단 완료 {cur['n_reports']}명이라 맞춤 체계도를 만들 수 없습니다(최소 {cur['min_reports']}명).")

    cid = ctx.company_id if not (ctx.is_super_admin and not ctx.is_acting) else company_id
    company = await db.get(Company, cid) if cid else None
    scope = company.name if company else "전체"
    if org_unit_id:
        lb = (await unit_labels(db, [org_unit_id])).get(org_unit_id) or {}
        scope = " › ".join(x for x in (scope, lb.get("division_name"), lb.get("office_name"), lb.get("team_name")) if x)

    NAVY, MID, BRIGHT, LIGHT, TINT, SOFT, LINE, GRAY = "1C2B4A", "2E4A85", "4F86F7", "8FB2FB", "EEF3FD", "F2F4F7", "D9DEE7", "64748B"
    fill = lambda c: PatternFill("solid", fgColor=c)
    thin = Side(style="thin", color=LINE)
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    tier_style = {
        "required": (fill(NAVY), Font(bold=True, color="FFFFFF", size=10), "필수"),
        "elective": (fill(TINT), Font(bold=True, color=NAVY, size=10), "선택"),
        "strength": (fill("E5EEFE"), Font(bold=True, color=MID, size=10), "강점 확산"),
    }
    tier_border = {
        "required": Border(left=Side(style="medium", color=NAVY), right=Side(style="medium", color=NAVY),
                           top=Side(style="medium", color=NAVY), bottom=Side(style="medium", color=NAVY)),
        "elective": Border(left=Side(style="thin", color=NAVY), right=Side(style="thin", color=NAVY),
                           top=Side(style="thin", color=NAVY), bottom=Side(style="thin", color=NAVY)),
        "strength": Border(left=Side(style="thin", color=BRIGHT), right=Side(style="thin", color=BRIGHT),
                           top=Side(style="thin", color=BRIGHT), bottom=Side(style="thin", color=BRIGHT)),
    }
    tracks = ["A", "B", "C", "D"]
    tmeta = {t["track"]: t for t in cur["tracks"]}
    tname = {"A": "기본기", "B": "실습", "C": "현업 적용", "D": "전수·확산"}

    wb = Workbook()
    ws = wb.active
    ws.title = "맞춤 교육체계도"
    ws.sheet_view.showGridLines = False
    last_col = 2 + len(tracks)

    # 제목 띠
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    c = ws.cell(row=1, column=1, value=f"{scope} 맞춤 리더십 교육체계")
    c.font = Font(bold=True, size=18, color="FFFFFF"); c.fill = fill(NAVY)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 40
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    c = ws.cell(row=2, column=1, value=cur["summary"])
    c.font = Font(size=10.5, color="FFFFFF"); c.fill = fill(NAVY); c.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[2].height = 30
    # 범례
    ws.cell(row=3, column=1, value=f"기준 {datetime.now():%Y.%m.%d} · 과정은 교육과정 Pool에서 진단 결과에 따라 자동 구성").font = Font(size=9, color=GRAY)
    lg = last_col - 2
    for i, t in enumerate(("required", "elective", "strength")):
        cell = ws.cell(row=3, column=lg + i, value=tier_style[t][2])
        cell.fill, cell.font, cell.alignment, cell.border = tier_style[t][0], tier_style[t][1], center, tier_border[t]
    ws.row_dimensions[3].height = 22

    # 머리 1단(남색 그룹 띠) + 2단(회색 트랙)
    hr1, hr2 = 5, 6
    # 왼쪽 두 머리 칸은 세로 병합 없이 두 줄을 같은 색으로 채운다(일부 미리보기에서 세로 병합 아래 칸이 밀리는 문제 회피)
    for col, txt in ((1, "역량"), (2, "하위역량")):
        top = ws.cell(row=hr1, column=col)
        top.fill = fill(NAVY)
        cell = ws.cell(row=hr2, column=col, value=txt)
        cell.fill, cell.font, cell.alignment = fill(NAVY), Font(bold=True, color="FFFFFF"), center
    ws.merge_cells(start_row=hr1, start_column=3, end_row=hr1, end_column=5)
    cell = ws.cell(row=hr1, column=3, value="역량 향상 과정 (개선)")
    cell.fill, cell.font, cell.alignment = fill(NAVY), Font(bold=True, color="FFFFFF", size=12), center
    cell = ws.cell(row=hr1, column=6, value="강점 확산 과정")
    cell.fill, cell.font, cell.alignment = fill(MID), Font(bold=True, color="FFFFFF", size=12), center
    for i, t in enumerate(tracks):
        m = tmeta.get(t, {})
        lv = f"Lv.{m.get('from_level')}→{m.get('to_level')}" if m.get("from_level") else ""
        cell = ws.cell(row=hr2, column=3 + i, value=f"{t} · {tname[t]}  {lv}\n{m.get('format') or ''}")
        cell.fill, cell.font, cell.alignment, cell.border = fill(SOFT), Font(bold=True, color=NAVY, size=10), center, box
    ws.row_dimensions[hr1].height = 26
    ws.row_dimensions[hr2].height = 36

    # 본문: 역량(세로 병합) × 하위역량 × 트랙
    r = hr2 + 1
    for ci, comp in enumerate(cur["competencies"]):
        rows = comp["rows"]
        if not rows:
            continue
        start = r
        for row in rows:
            dist = row["dist"]
            sub_cell = ws.cell(row=r, column=2, value=f"{row['sub']}\n측정 {row['measured']}명 · Lv1 {dist['1']} · Lv2 {dist['2']} · Lv3 {dist['3']} · Lv4 {dist['4']}")
            sub_cell.fill, sub_cell.font, sub_cell.border = fill("FAFBFC"), Font(bold=True, size=10, color="0F172A"), box
            sub_cell.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
            by_track = {i["track"]: i for i in row["items"]}
            for ti, t in enumerate(tracks):
                cell = ws.cell(row=r, column=3 + ti)
                it = by_track.get(t)
                if it:
                    f, fnt, label = tier_style[it["tier"]]
                    cell.value = f"[{label}] {it['course']}\n{it['people']}명"
                    cell.fill, cell.font, cell.border, cell.alignment = f, fnt, tier_border[it["tier"]], center
                else:
                    cell.border = box
            ws.row_dimensions[r].height = 46
            r += 1
        ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
        head = comp["name"] + ("\n집중 역량" if comp["priority"] == "focus" else "")
        head += f"\n평균 {comp['avg']:.1f}" if comp["avg"] is not None else ""
        if comp.get("baseline_avg") is not None:
            head += f"\n회사 {comp['baseline_avg']:.1f}"
        cell = ws.cell(row=start, column=1, value=head)
        cell.fill = fill(NAVY if ci % 2 == 0 else MID)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = center
        for rr in range(start, r):
            ws.cell(row=rr, column=1).border = box

    # 근거 시트(필터·정렬용 평평한 표)
    ws2 = wb.create_sheet("구성 근거")
    ws2.append(["역량", "집중 역량", "하위역량", "구분", "트랙", "과정", "교육목표", "대상 인원", "근거"])
    for cell in ws2[1]:
        cell.fill, cell.font = fill(NAVY), Font(bold=True, color="FFFFFF")
    for comp in cur["competencies"]:
        for row in comp["rows"]:
            for it in row["items"]:
                ws2.append([comp["name"], "예" if comp["priority"] == "focus" else "", row["sub"],
                            tier_style[it["tier"]][2], it["track"], it["course"], it["goal"], it["people"], it["reason"]])
    for col, w in zip("ABCDEFGHI", (10, 8, 18, 9, 6, 26, 40, 9, 46)):
        ws2.column_dimensions[col].width = w

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 30
    for i in range(len(tracks)):
        ws.column_dimensions[get_column_letter(3 + i)].width = 28
    ws.freeze_panes = ws.cell(row=hr2 + 1, column=3)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"{scope.replace(' › ', '_')}_맞춤교육체계도_{datetime.now():%Y%m%d}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=curriculum.xlsx; filename*=UTF-8''{quote(fname)}"},
    )
