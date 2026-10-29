# diag_project/routes/roster.py
#
# 참여자 명단(엑셀) → 조직 구조 + 개인 참여 코드 자동 세팅.
#   - 업로드한 명단의 본부·실·팀 이름으로 조직 단위를 만들고(없을 때만, 코드 자동),
#     참여자마다 개인 참여 코드(예: AS1-K7Q2)를 발급한다.
#   - 결과는 그 고객사 HR 화면과 마스터 화면에 그대로 나타난다(같은 데이터).
#   - 권한: 마스터는 모든 고객사, 고객사 HR 은 자사만. 미리보기(dry_run) 후 확정.
#   - 한 행이라도 문제가 있으면 아무것도 저장하지 않고 행별 사유를 돌려준다(부분 등록 방지).

import io
import os
import re
import secrets
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import quote
from uuid import UUID

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select

from diag_project.database import get_db
from diag_project.models.company import Company
from diag_project.models.org_unit import OrgUnit
from diag_project.models.participant import Participant
from diag_project.services.auth import AdminContext, get_current_admin

router = APIRouter(prefix="/admin", tags=["Admin · 참여자 명단"])

# 헷갈리는 글자(I, O, 0, 1) 없는 4자 → 개인 코드 뒷부분
_CODE_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
MAX_ROWS = 3000

# 머리글 인식(공백 무시). 첫 번째로 맞는 열을 쓴다.
_HEADERS = {
    "name": ("이름", "성명", "name"),
    "email": ("이메일", "메일", "email", "e-mail", "이메일주소"),
    "division": ("본부", "부문", "division"),
    "office": ("실", "office"),
    "team": ("팀", "부서", "team", "department"),
    "employee_no": ("사번", "직원번호", "employeeno"),
}


def _norm_header(v: Any) -> str:
    return re.sub(r"\s+", "", str(v or "")).lower()


def parse_roster(content: bytes) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """엑셀 바이트 → (행 목록, 머리글 오류). 머리글은 위 10행 안에서 '이름'과 '이메일'이 있는 행."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return [], [{"row": 0, "reason": "엑셀(.xlsx) 파일을 읽지 못했습니다."}]
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header_idx, colmap = None, {}
    for i, r in enumerate(rows[:10]):
        hs = [_norm_header(c) for c in r]
        cm = {}
        for key, names in _HEADERS.items():
            for j, h in enumerate(hs):
                if h in names and key not in cm:
                    cm[key] = j
        if "name" in cm and "email" in cm:
            header_idx, colmap = i, cm
            break
    if header_idx is None:
        return [], [{"row": 0, "reason": "머리글에 '이름'과 '이메일' 열이 필요합니다(본부·실·팀은 선택)."}]

    out = []
    for k, r in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        def cell(key: str) -> str:
            j = colmap.get(key)
            return str(r[j]).strip() if (j is not None and j < len(r) and r[j] is not None) else ""
        if not any(cell(x) for x in ("name", "email", "division", "office", "team")):
            continue
        out.append({"row": k, "name": cell("name"), "email": cell("email").lower(),
                    "division": cell("division"), "office": cell("office"), "team": cell("team"),
                    "employee_no": cell("employee_no")})
    return out, []


async def _unique_access_code(db: AsyncSession, base: str, taken: set) -> str:
    for _ in range(50):
        code = f"{base}-{''.join(secrets.choice(_CODE_ALPHABET) for _ in range(4))}"
        if code in taken:
            continue
        if (await db.execute(select(Participant.id).where(Participant.access_code == code))).first():
            continue
        taken.add(code)
        return code
    raise HTTPException(status_code=500, detail="개인 코드를 만들지 못했습니다. 다시 시도해 주세요.")


async def _find_or_create_units(db: AsyncSession, company: Company, row: Dict[str, Any],
                                cache: Dict[Tuple, OrgUnit], created: List[dict]) -> Optional[OrgUnit]:
    """명단 한 행의 본부 › 실 › 팀 이름 → 가장 아래 조직 단위(없으면 자동 코드로 만든다)."""
    from diag_project.routes.org_units import _create_division, _create_office, _create_team

    div_name, off_name, team_name = row["division"], row["office"], row["team"]
    if not div_name:
        if off_name or team_name:
            raise ValueError("실·팀이 있으면 본부도 적어 주세요.")
        return None

    async def get(level: str, name: str, parent: Optional[OrgUnit]) -> Optional[OrgUnit]:
        key = (level, name, parent.id if parent else None)
        if key in cache:
            return cache[key]
        q = select(OrgUnit).where(OrgUnit.company_id == company.id, OrgUnit.level == level, OrgUnit.name == name)
        q = q.where(OrgUnit.parent_id == parent.id) if parent else q.where(OrgUnit.parent_id.is_(None))
        u = (await db.execute(q)).scalars().first()
        if u is not None:
            cache[key] = u
        return u

    async def ensure(level: str, name: str, parent: Optional[OrgUnit]) -> OrgUnit:
        u = await get(level, name, parent)
        if u is None:
            try:
                if level == "division":
                    u = await _create_division(db, company, name, "")
                elif level == "office":
                    u = await _create_office(db, company, parent, name, "")
                else:
                    u = await _create_team(db, company, parent, name, None)
            except HTTPException as e:
                raise ValueError(str(e.detail))
            cache[(level, name, parent.id if parent else None)] = u
            created.append({"code": u.code, "name": u.name, "level": u.level})
        if not u.is_active:
            raise ValueError(f"'{u.name}'({u.code})은 사용 중지된 조직입니다.")
        return u

    unit = await ensure("division", div_name, None)
    if off_name:
        unit = await ensure("office", off_name, unit)
    if team_name:
        unit = await ensure("team", team_name, unit)
    return unit


@router.post("/companies/{company_id}/roster")
async def import_roster(
    company_id: UUID,
    file: UploadFile = File(...),
    dry_run: bool = Query(True, description="true 면 저장하지 않고 미리보기만"),
    ctx: AdminContext = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db),
):
    """명단 엑셀(이름·이메일·본부·실·팀) → 조직 자동 세팅 + 참여자 등록 + 개인 코드 발급."""
    ctx.assert_can_access_company(company_id)
    company = await db.get(Company, company_id)
    if company is None:
        raise HTTPException(status_code=404, detail="고객사를 찾을 수 없습니다.")
    content = await file.read()
    rows, header_errors = parse_roster(content)
    if header_errors:
        raise HTTPException(status_code=400, detail=header_errors[0]["reason"])
    if not rows:
        raise HTTPException(status_code=400, detail="명단에 참여자가 없습니다.")
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"한 번에 {MAX_ROWS}명까지 올릴 수 있습니다.")

    errors: List[Dict[str, Any]] = []
    seen_emails: Dict[str, int] = {}
    for r in rows:
        if not r["name"]:
            errors.append({"row": r["row"], "reason": "이름이 없습니다."})
        if not _EMAIL_RE.match(r["email"]):
            errors.append({"row": r["row"], "reason": f"이메일 형식이 올바르지 않습니다({r['email'] or '빈 칸'})."})
        elif r["email"] in seen_emails:
            errors.append({"row": r["row"], "reason": f"같은 이메일이 {seen_emails[r['email']]}행에도 있습니다."})
        else:
            seen_emails[r["email"]] = r["row"]

    existing = {p.email.lower(): p for p in (await db.execute(
        select(Participant).where(Participant.email.in_(list(seen_emails)))
    )).scalars().all()} if seen_emails else {}
    for r in rows:
        p = existing.get(r["email"])
        if p is not None and p.company_id and p.company_id != company.id:
            errors.append({"row": r["row"], "reason": "다른 고객사에 이미 등록된 이메일입니다."})

    created_units: List[dict] = []
    preview: List[Dict[str, Any]] = []
    new_count = upd_count = 0
    cache: Dict[Tuple, OrgUnit] = {}
    taken: set = set()
    try:
        for r in rows:
            try:
                unit = await _find_or_create_units(db, company, r, cache, created_units)
            except ValueError as e:
                errors.append({"row": r["row"], "reason": str(e)})
                continue
            if any(e["row"] == r["row"] for e in errors):
                continue
            p = existing.get(r["email"])
            base = unit.code if unit else (company.org_prefix or "")
            if not base:
                from diag_project.routes.org_units import _ensure_prefix
                base = await _ensure_prefix(db, company)
            if p is None:
                p = Participant(email=r["email"], name=r["name"], group_code=company.code, company_id=company.id,
                                org_unit_id=unit.id if unit else None, password_hash="dummy_hashed_password",
                                access_code=await _unique_access_code(db, base, taken))
                db.add(p)
                new_count += 1
                action = "신규"
            else:
                p.name = r["name"] or p.name
                p.company_id = company.id
                p.group_code = company.code
                p.org_unit_id = unit.id if unit else p.org_unit_id
                # 소속이 바뀌면 코드도 새 소속 기준으로 다시 만든다(이전 코드는 더 이상 안 통함)
                if not p.access_code or (unit and not p.access_code.startswith(unit.code + "-")):
                    p.access_code = await _unique_access_code(db, base, taken)
                db.add(p)
                upd_count += 1
                action = "갱신"
            preview.append({"row": r["row"], "name": r["name"], "email": r["email"], "action": action,
                            "org_code": unit.code if unit else None,
                            "org_path": " › ".join(x for x in (r["division"], r["office"], r["team"]) if x) or "미분류"})
        await db.flush()
    except HTTPException:
        await db.rollback()
        raise

    ok = not errors
    if dry_run or not ok:
        await db.rollback()
    else:
        await db.commit()
    return {
        "dry_run": dry_run, "saved": (ok and not dry_run),
        "rows": len(rows), "new": new_count, "updated": upd_count,
        "org_units_created": created_units,
        "errors": sorted(errors, key=lambda e: e["row"]),
        "preview": preview[:200],
    }


def _xlsx_response(wb, fname: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=roster.xlsx; filename*=UTF-8''{quote(fname)}"},
    )


@router.get("/roster/template")
async def roster_template(ctx: AdminContext = Depends(get_current_admin)):
    """명단 양식(머리글 + 예시 2행). 본부·실·팀은 비워도 된다."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "참여자 명단"
    ws.append(["이름", "이메일", "본부", "실", "팀", "사번"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1C2B4A")
    ws.append(["홍길동", "gildong@company.com", "영업본부", "", "영업1팀", "1001"])
    ws.append(["김하늘", "haneul@company.com", "영업본부", "마케팅실", "브랜드팀", "1002"])
    for col, w in zip("ABCDEF", (12, 28, 16, 16, 16, 10)):
        ws.column_dimensions[col].width = w
    return _xlsx_response(wb, "참여자_명단_양식.xlsx")


@router.get("/companies/{company_id}/roster/export")
async def export_roster_codes(
    company_id: UUID,
    ctx: AdminContext = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db),
):
    """참여자별 개인 참여 코드 엑셀(고객사 전달용). 소속순 → 이름순. 진단 상태 포함."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from diag_project.services.org_units import participant_statuses, unit_labels

    ctx.assert_can_access_company(company_id)
    company = await db.get(Company, company_id)
    if company is None:
        raise HTTPException(status_code=404, detail="고객사를 찾을 수 없습니다.")
    ps = (await db.execute(select(Participant).where(Participant.company_id == company.id))).scalars().all()
    if not ps:
        raise HTTPException(status_code=400, detail="등록된 참여자가 없습니다. 먼저 명단을 올려 주세요.")
    labels = await unit_labels(db, [p.org_unit_id for p in ps])
    st = await participant_statuses(db, [p.id for p in ps])
    st_ko = {"completed": "완료", "in_progress": "진행 중", "abandoned": "중단", "not_started": "미시작"}

    def key(p: Participant):
        lb = labels.get(p.org_unit_id) or {}
        return (lb.get("code") is None, lb.get("code") or "", p.name or "")

    navy = PatternFill("solid", fgColor="1C2B4A")
    tint = PatternFill("solid", fgColor="EEF3FD")
    thin = Side(style="thin", color="D9DEE7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = Workbook()
    ws = wb.active
    ws.title = "개인 참여 코드"
    ws["A1"] = f"{company.name} 리더십 진단 개인 참여 코드"
    ws["A1"].font = Font(size=15, bold=True, color="1C2B4A")
    ws["A2"] = f"접속 주소: {os.getenv('FM_PARTICIPANT_URL', 'https://fm.connectn.co.kr')}"
    ws["A3"] = "로그인 화면에 이름·이메일과 함께 '그룹 코드' 칸에 개인 참여 코드를 입력하세요. 이메일이 명단과 같아야 들어갈 수 있습니다."
    if company.diagnosis_start and company.diagnosis_end:
        ws["A4"] = f"진단 기간: {company.diagnosis_start:%Y.%m.%d} ~ {company.diagnosis_end:%Y.%m.%d}"
    head = ["이름", "이메일", "본부", "실", "팀", "개인 참여 코드", "진단 상태"]
    for col, h in enumerate(head, start=1):
        c = ws.cell(row=6, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = navy
        c.alignment = Alignment(horizontal="center")
        c.border = border
    r = 7
    for p in sorted(ps, key=key):
        lb = labels.get(p.org_unit_id) or {}
        vals = [p.name, p.email, lb.get("division_name") or "", lb.get("office_name") or "", lb.get("team_name") or "",
                p.access_code or "(명단 외 참여)", st_ko.get(st.get(p.id, "not_started"), "")]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = border
            if col == 6:
                c.font = Font(bold=True, color="1C2B4A")
                c.fill = tint
        r += 1
    for col, w in zip("ABCDEFG", (12, 30, 16, 16, 16, 18, 10)):
        ws.column_dimensions[col].width = w
    return _xlsx_response(wb, f"{company.name}_개인참여코드_{datetime.now():%Y%m%d}.xlsx")
