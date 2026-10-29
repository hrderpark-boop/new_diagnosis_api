# diag_project/routes/org_units.py
#
# 조직 참여 코드(AS1 형식) 관리 API.
#   - 마스터(super_admin): 고객사 접두어 지정, 본부·실·팀 등록·수정·삭제, 일괄 등록
#   - 고객사 HR(client_admin): 자사 조직 트리·조직별 통계 조회, 자사 진단 기간 설정
#   - 진단 기간: 마스터·HR 모두 설정(HR 은 자사만)
# 코드 규칙은 services/org_codes.py, 트리·집계는 services/org_units.py.

from datetime import date
from typing import List, Optional
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, status
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select

from diag_project.database import get_db
from diag_project.models.company import Company
from diag_project.models.org_unit import OrgUnit
from diag_project.models.participant import Participant
from diag_project.services import org_codes
from diag_project.services.auth import AdminContext, get_current_admin, require_super_admin
from diag_project.services.org_units import build_org_tree, code_in_use

router = APIRouter(prefix="/admin", tags=["Admin · 조직 코드"])


# ---------------------------------------------------------------------------
# 스키마
# ---------------------------------------------------------------------------
class CompanyUpdateRequest(BaseModel):
    name: Optional[str] = None
    contact_email: Optional[str] = None
    is_active: Optional[bool] = None
    org_prefix: Optional[str] = Field(None, description="조직 코드 접두어(영문 1~3자)")


class OrgUnitCreateRequest(BaseModel):
    level: str = Field(..., description="division | office | team")
    name: str
    # 본부·실: 영문 1~2자 코드(필수). 팀: 무시.
    short_code: Optional[str] = None
    # 실: 소속 본부 id(필수). 팀: 소속 본부 또는 실 id(필수)와 번호(비우면 다음 번호)
    parent_id: Optional[UUID] = None
    number: Optional[int] = None


class OrgUnitUpdateRequest(BaseModel):
    name: Optional[str] = None
    is_active: Optional[bool] = None
    sort_order: Optional[int] = None


class OrgBulkRow(BaseModel):
    division: str
    division_code: Optional[str] = None   # 비우면 자동
    office: Optional[str] = None
    office_code: Optional[str] = None
    team: Optional[str] = None


class PeriodRequest(BaseModel):
    diagnosis_start: Optional[date] = None
    diagnosis_end: Optional[date] = None


class OrgBulkRequest(BaseModel):
    rows: List[OrgBulkRow]


# ---------------------------------------------------------------------------
# 내부 도우미
# ---------------------------------------------------------------------------
async def _get_company(db: AsyncSession, company_id: UUID) -> Company:
    company = await db.get(Company, company_id)
    if company is None:
        raise HTTPException(status_code=404, detail="고객사를 찾을 수 없습니다.")
    return company


async def _ensure_prefix(db: AsyncSession, company: Company) -> str:
    """접두어가 없으면 다른 고객사와 겹치지 않는 가장 짧은 접두어를 자동으로 정한다(같은 트랜잭션)."""
    if company.org_prefix:
        return company.org_prefix
    others = (await db.execute(
        select(Company.org_prefix).where(Company.id != company.id, Company.org_prefix.is_not(None))
    )).scalars().all()
    company.org_prefix = org_codes.auto_prefix(others, hints=[company.code or "", company.name or ""])
    db.add(company)
    await db.flush()
    return company.org_prefix


async def _auto_letter(db: AsyncSession, company_id: UUID, parent_id: Optional[UUID], level: str, base_code: str) -> str:
    """같은 상위 안에서 아직 안 쓴 다음 영문 코드. 완성 코드가 다른 곳과 겹치면 다음 글자로."""
    q = select(OrgUnit.short_code).where(OrgUnit.company_id == company_id, OrgUnit.level == level)
    q = q.where(OrgUnit.parent_id == parent_id) if parent_id else q.where(OrgUnit.parent_id.is_(None))
    used = set((await db.execute(q)).scalars().all())
    while True:
        short = org_codes.next_letter_code(used)
        if not await code_in_use(db, base_code + short):
            return short
        used.add(short)


async def _create_division(db: AsyncSession, company: Company, name: str, short: str) -> OrgUnit:
    prefix = await _ensure_prefix(db, company)
    if not (short or "").strip():   # 비우면 자동(A, B, C …)
        short = await _auto_letter(db, company.id, None, "division", prefix)
    err = org_codes.validate_division_code(short)
    if err:
        raise HTTPException(status_code=400, detail=err)
    code = org_codes.division_code(prefix, short)
    if await code_in_use(db, code):
        raise HTTPException(status_code=409, detail=f"이미 쓰는 코드입니다({code}).")
    unit = OrgUnit(company_id=company.id, level="division", name=name.strip(),
                   short_code=org_codes.normalize(short), code=code)
    db.add(unit)
    await db.flush()
    return unit


async def _create_office(db: AsyncSession, company: Company, parent: OrgUnit, name: str, short: str) -> OrgUnit:
    if parent.company_id != company.id or parent.level != "division":
        raise HTTPException(status_code=400, detail="실은 같은 고객사의 본부 아래에만 만들 수 있습니다.")
    if not (short or "").strip():   # 비우면 자동
        short = await _auto_letter(db, company.id, parent.id, "office", parent.code)
    err = org_codes.validate_office_code(short)
    if err:
        raise HTTPException(status_code=400, detail=err)
    code = org_codes.office_code(parent.code, short)
    if await code_in_use(db, code):
        raise HTTPException(status_code=409, detail=f"이미 쓰는 코드입니다({code}).")
    unit = OrgUnit(company_id=company.id, parent_id=parent.id, level="office", name=name.strip(),
                   short_code=org_codes.normalize(short), code=code)
    db.add(unit)
    await db.flush()
    return unit


async def _create_team(db: AsyncSession, company: Company, parent: OrgUnit, name: str,
                       number: Optional[int]) -> OrgUnit:
    if parent.company_id != company.id or parent.level not in ("division", "office"):
        raise HTTPException(status_code=400, detail="팀은 같은 고객사의 본부나 실 아래에만 만들 수 있습니다.")
    if number is None:
        existing = (await db.execute(
            select(OrgUnit.short_code).where(OrgUnit.parent_id == parent.id)
        )).scalars().all()
        number = org_codes.next_team_number(existing)
    err = org_codes.validate_team_number(number)
    if err:
        raise HTTPException(status_code=400, detail=err)
    code = org_codes.team_code(parent.code, number)
    if await code_in_use(db, code):
        raise HTTPException(status_code=409, detail=f"이미 쓰는 코드입니다({code}).")
    unit = OrgUnit(company_id=company.id, parent_id=parent.id, level="team", name=name.strip(),
                   short_code=str(int(number)), code=code, sort_order=int(number))
    db.add(unit)
    await db.flush()
    return unit


# ---------------------------------------------------------------------------
# 고객사 수정(접두어 포함) — 마스터 전용
# ---------------------------------------------------------------------------
@router.patch("/companies/{company_id}")
async def update_company(
    company_id: UUID,
    body: CompanyUpdateRequest,
    ctx: AdminContext = Depends(require_super_admin),
    db: AsyncSession = Depends(get_db),
):
    company = await _get_company(db, company_id)
    if body.name is not None:
        if not body.name.strip():
            raise HTTPException(status_code=400, detail="고객사 이름은 비울 수 없습니다.")
        company.name = body.name.strip()
    if body.contact_email is not None:
        company.contact_email = body.contact_email.strip() or None
    if body.is_active is not None:
        company.is_active = body.is_active
    if body.org_prefix is not None:
        new_prefix = org_codes.normalize(body.org_prefix)
        if new_prefix != (company.org_prefix or ""):
            has_units = (await db.execute(
                select(OrgUnit.id).where(OrgUnit.company_id == company.id)
            )).first()
            if has_units:
                raise HTTPException(status_code=409,
                                    detail="조직 코드가 이미 있어 접두어를 바꿀 수 없습니다(배포한 코드가 깨집니다).")
            others = (await db.execute(
                select(Company.org_prefix).where(Company.id != company.id, Company.org_prefix.is_not(None))
            )).scalars().all()
            err = org_codes.validate_prefix(new_prefix, others)
            if err:
                raise HTTPException(status_code=400, detail=err)
            company.org_prefix = new_prefix
    db.add(company)
    await db.commit()
    await db.refresh(company)
    return {"id": str(company.id), "name": company.name, "code": company.code,
            "contact_email": company.contact_email, "is_active": company.is_active,
            "org_prefix": company.org_prefix}


# ---------------------------------------------------------------------------
# 조직 트리 조회 — 마스터는 모든 고객사, HR 은 자사만
# ---------------------------------------------------------------------------
@router.get("/companies/{company_id}/org-units")
async def get_org_units(
    company_id: UUID,
    ctx: AdminContext = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db),
):
    ctx.assert_can_access_company(company_id)
    company = await _get_company(db, company_id)
    tree = await build_org_tree(db, company.id)
    return {
        "company": {"id": str(company.id), "name": company.name, "code": company.code,
                    "org_prefix": company.org_prefix},
        **tree,
    }


# ---------------------------------------------------------------------------
# 본부·팀 등록 — 마스터 전용
# ---------------------------------------------------------------------------
@router.post("/companies/{company_id}/org-units", status_code=status.HTTP_201_CREATED)
async def create_org_unit(
    company_id: UUID,
    body: OrgUnitCreateRequest,
    ctx: AdminContext = Depends(require_super_admin),
    db: AsyncSession = Depends(get_db),
):
    company = await _get_company(db, company_id)
    if not body.name or not body.name.strip():
        raise HTTPException(status_code=400, detail="이름을 입력해 주세요.")
    if body.level == "division":
        unit = await _create_division(db, company, body.name, body.short_code or "")
    elif body.level in ("office", "team"):
        if not body.parent_id:
            raise HTTPException(status_code=400, detail="소속 본부(또는 실)를 골라야 합니다.")
        parent = await db.get(OrgUnit, body.parent_id)
        if parent is None:
            raise HTTPException(status_code=404, detail="상위 조직을 찾을 수 없습니다.")
        if body.level == "office":
            unit = await _create_office(db, company, parent, body.name, body.short_code or "")
        else:
            unit = await _create_team(db, company, parent, body.name, body.number)
    else:
        raise HTTPException(status_code=400, detail="level 은 division, office, team 중 하나입니다.")
    await db.commit()
    return {"id": str(unit.id), "code": unit.code, "name": unit.name, "level": unit.level,
            "parent_id": str(unit.parent_id) if unit.parent_id else None}


@router.post("/companies/{company_id}/org-units/bulk", status_code=status.HTTP_201_CREATED)
async def bulk_create_org_units(
    company_id: UUID,
    body: OrgBulkRequest,
    ctx: AdminContext = Depends(require_super_admin),
    db: AsyncSession = Depends(get_db),
):
    """엑셀 붙여넣기용 일괄 등록. 행 = (본부, 본부 코드, [실, 실 코드], 팀). 이미 있는 본부·실·팀은 건너뛴다.

    한 행이라도 규칙 위반이면 전부 되돌린다(부분 등록 방지).
    """
    company = await _get_company(db, company_id)
    prefix = await _ensure_prefix(db, company)
    created: List[dict] = []
    skipped = 0
    existing_divs = (await db.execute(
        select(OrgUnit).where(OrgUnit.company_id == company.id, OrgUnit.level == "division")
    )).scalars().all()
    divisions = {u.short_code: u for u in existing_divs}
    divisions_by_name = {u.name: u for u in existing_divs}
    try:
        for i, row in enumerate(body.rows, start=1):
            short = org_codes.normalize(row.division_code)
            # 코드를 비우면 같은 이름의 본부를 쓰고, 없으면 자동 코드로 새로 만든다
            div = divisions.get(short) if short else divisions_by_name.get(row.division.strip())
            if div is None:
                try:
                    div = await _create_division(db, company, row.division, short)
                except HTTPException as e:
                    raise HTTPException(status_code=e.status_code, detail=f"{i}행: {e.detail}")
                divisions[div.short_code] = div
                divisions_by_name[div.name] = div
                created.append({"code": div.code, "name": div.name, "level": "division"})
            elif div.name != row.division.strip():
                raise HTTPException(status_code=400,
                                    detail=f"{i}행: 본부 코드 {div.code} 는 이미 '{div.name}' 입니다.")
            parent = div
            office_name = (row.office or "").strip()
            if office_name:
                oshort = org_codes.normalize(row.office_code)
                oq = select(OrgUnit).where(OrgUnit.parent_id == div.id, OrgUnit.level == "office")
                oq = oq.where(OrgUnit.short_code == oshort) if oshort else oq.where(OrgUnit.name == office_name)
                off = (await db.execute(oq)).scalars().first()
                if off is None:
                    try:
                        off = await _create_office(db, company, div, office_name, oshort)
                    except HTTPException as e:
                        raise HTTPException(status_code=e.status_code, detail=f"{i}행: {e.detail}")
                    created.append({"code": off.code, "name": off.name, "level": "office"})
                elif off.name != office_name:
                    raise HTTPException(status_code=400,
                                        detail=f"{i}행: 실 코드 {off.code} 는 이미 '{off.name}' 입니다.")
                parent = off
            team_name = (row.team or "").strip()
            if not team_name:
                continue
            same = (await db.execute(
                select(OrgUnit.id).where(OrgUnit.parent_id == parent.id, OrgUnit.name == team_name)
            )).first()
            if same:
                skipped += 1
                continue
            try:
                t = await _create_team(db, company, parent, team_name, None)
            except HTTPException as e:
                raise HTTPException(status_code=e.status_code, detail=f"{i}행: {e.detail}")
            created.append({"code": t.code, "name": t.name, "level": "team"})
        await db.commit()
    except HTTPException:
        await db.rollback()
        raise
    return {"prefix": prefix, "created": created, "skipped": skipped}


# ---------------------------------------------------------------------------
# 단위 수정·삭제 — 마스터 전용
# ---------------------------------------------------------------------------
@router.patch("/org-units/{unit_id}")
async def update_org_unit(
    unit_id: UUID,
    body: OrgUnitUpdateRequest,
    ctx: AdminContext = Depends(require_super_admin),
    db: AsyncSession = Depends(get_db),
):
    """이름·사용 여부·순서만 바꾼다. 코드는 이미 배포됐을 수 있어 바꾸지 않는다."""
    unit = await db.get(OrgUnit, unit_id)
    if unit is None:
        raise HTTPException(status_code=404, detail="조직을 찾을 수 없습니다.")
    if body.name is not None:
        if not body.name.strip():
            raise HTTPException(status_code=400, detail="이름은 비울 수 없습니다.")
        unit.name = body.name.strip()
    if body.is_active is not None:
        unit.is_active = body.is_active
    if body.sort_order is not None:
        unit.sort_order = body.sort_order
    db.add(unit)
    await db.commit()
    return {"id": str(unit.id), "code": unit.code, "name": unit.name, "is_active": unit.is_active}


@router.delete("/org-units/{unit_id}")
async def delete_org_unit(
    unit_id: UUID,
    ctx: AdminContext = Depends(require_super_admin),
    db: AsyncSession = Depends(get_db),
):
    """하위 팀이나 소속 참여자가 있으면 지우지 않는다(사용 중지로 막으면 된다)."""
    unit = await db.get(OrgUnit, unit_id)
    if unit is None:
        raise HTTPException(status_code=404, detail="조직을 찾을 수 없습니다.")
    if (await db.execute(select(OrgUnit.id).where(OrgUnit.parent_id == unit.id))).first():
        raise HTTPException(status_code=409, detail="소속 팀이 있어 삭제할 수 없습니다. 팀을 먼저 지우거나 사용 중지하세요.")
    if (await db.execute(select(Participant.id).where(Participant.org_unit_id == unit.id))).first():
        raise HTTPException(status_code=409, detail="이 코드로 참여한 사람이 있어 삭제할 수 없습니다. 사용 중지하세요.")
    await db.delete(unit)
    await db.commit()
    return {"deleted": True}



# ---------------------------------------------------------------------------
# 진단 기간 — 마스터는 모든 고객사, HR 은 자사만
# ---------------------------------------------------------------------------
def _period_out(c: Company) -> dict:
    return {"company_id": str(c.id), "diagnosis_start": c.diagnosis_start, "diagnosis_end": c.diagnosis_end}


@router.get("/companies/{company_id}/period")
async def get_period(
    company_id: UUID,
    ctx: AdminContext = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db),
):
    ctx.assert_can_access_company(company_id)
    return _period_out(await _get_company(db, company_id))


@router.put("/companies/{company_id}/period")
async def set_period(
    company_id: UUID,
    body: PeriodRequest,
    ctx: AdminContext = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db),
):
    """시작·종료일을 함께 정하거나(둘 다) 함께 지운다(둘 다 null). 최대 1년."""
    ctx.assert_can_access_company(company_id)
    company = await _get_company(db, company_id)
    s, e = body.diagnosis_start, body.diagnosis_end
    if (s is None) != (e is None):
        raise HTTPException(status_code=400, detail="시작일과 종료일을 함께 정해 주세요.")
    if s and e:
        if s > e:
            raise HTTPException(status_code=400, detail="종료일이 시작일보다 빠릅니다.")
        if (e - s).days > 366:
            raise HTTPException(status_code=400, detail="진단 기간은 1년을 넘을 수 없습니다.")
    company.diagnosis_start, company.diagnosis_end = s, e
    db.add(company)
    await db.commit()
    await db.refresh(company)
    return _period_out(company)


# ---------------------------------------------------------------------------
# 조직별 진행 현황 · 역량 평균 — 마스터는 모든 고객사, HR 은 자사만
# ---------------------------------------------------------------------------
@router.get("/companies/{company_id}/org-stats")
async def get_org_stats(
    company_id: UUID,
    level: str = "division",
    ctx: AdminContext = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db),
):
    if level not in ("division", "office", "team"):
        raise HTTPException(status_code=400, detail="level 은 division, office, team 중 하나입니다.")
    ctx.assert_can_access_company(company_id)
    await _get_company(db, company_id)
    from diag_project.data.competencies import COMPETENCY_FRAMEWORK
    from diag_project.services.org_units import org_stats
    order = list(COMPETENCY_FRAMEWORK.keys())
    out = await org_stats(db, company_id, level, order)
    out["competencies"] = [{"key": k, "name": COMPETENCY_FRAMEWORK[k].get("name", k)} for k in order]
    return out



# ---------------------------------------------------------------------------
# 참여 코드 엑셀 — 고객사 전달용(마스터는 모든 고객사, HR 은 자사만)
# ---------------------------------------------------------------------------
@router.get("/companies/{company_id}/org-codes/export")
async def export_org_codes(
    company_id: UUID,
    ctx: AdminContext = Depends(get_current_admin),
    db: AsyncSession = Depends(get_db),
):
    """조직별 참여 코드 엑셀. 1시트 = 배포용(팀 단위, 팀이 없는 실·본부는 그 단위), 2시트 = 전체 코드 목록.

    사용 중지한 코드는 배포용 시트에서 뺀다(전체 목록에는 '사용 중지'로 남김).
    """
    import io
    import os
    from datetime import datetime
    from urllib.parse import quote

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    ctx.assert_can_access_company(company_id)
    company = await _get_company(db, company_id)
    units = (await db.execute(
        select(OrgUnit).where(OrgUnit.company_id == company.id).order_by(OrgUnit.code)
    )).scalars().all()
    if not units:
        raise HTTPException(status_code=400, detail="등록된 조직 코드가 없습니다. 먼저 본부·실·팀을 등록해 주세요.")
    by_id = {u.id: u for u in units}
    children = {}
    for u in units:
        if u.parent_id:
            children.setdefault(u.parent_id, []).append(u)

    def chain(u: OrgUnit):
        out, cur = {}, u
        while cur is not None:
            out[cur.level] = cur
            cur = by_id.get(cur.parent_id) if cur.parent_id else None
        return out

    login_url = os.getenv("FM_PARTICIPANT_URL", "https://fm.connectn.co.kr")
    navy = PatternFill("solid", fgColor="1C2B4A")
    tint = PatternFill("solid", fgColor="EEF3FD")
    thin = Side(style="thin", color="D9DEE7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "참여 코드(배포용)"
    ws["A1"] = f"{company.name} 리더십 진단 참여 코드"
    ws["A1"].font = Font(size=15, bold=True, color="1C2B4A")
    ws["A2"] = f"접속 주소: {login_url}"
    ws["A3"] = "로그인 화면의 '그룹 코드' 칸에 본인 소속의 참여 코드를 입력하세요. 코드로 회사와 소속이 자동으로 정해집니다."
    if company.diagnosis_start and company.diagnosis_end:
        ws["A4"] = f"진단 기간: {company.diagnosis_start:%Y.%m.%d} ~ {company.diagnosis_end:%Y.%m.%d}"
    headers = ["본부", "실", "팀", "참여 코드"]
    hr = 6
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=hr, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = navy
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    r = hr + 1
    for u in units:
        if not u.is_active or any(ch.is_active for ch in children.get(u.id, [])):
            continue   # 배포는 가장 아래 단위(사용 중) 코드만
        ch = chain(u)
        vals = [ch.get("division").name if ch.get("division") else "",
                ch.get("office").name if ch.get("office") else "",
                ch.get("team").name if ch.get("team") else "", u.code]
        for col, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = border
            if col == 4:
                c.font = Font(bold=True, color="1C2B4A", size=12)
                c.fill = tint
                c.alignment = Alignment(horizontal="center")
        r += 1
    for col, w in zip("ABCD", (22, 22, 22, 16)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("전체 코드")
    lv = {"division": "본부", "office": "실", "team": "팀"}
    ws2.append(["구분", "본부", "실", "팀", "참여 코드", "상태"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = navy
    for u in units:
        ch = chain(u)
        ws2.append([lv.get(u.level, u.level),
                    ch.get("division").name if ch.get("division") else "",
                    ch.get("office").name if ch.get("office") else "",
                    ch.get("team").name if ch.get("team") else "",
                    u.code, "사용" if u.is_active else "사용 중지"])
    for col, w in zip("ABCDEF", (8, 22, 22, 22, 14, 10)):
        ws2.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"{company.name}_참여코드_{datetime.now():%Y%m%d}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=org_codes.xlsx; filename*=UTF-8''{quote(fname)}"},
    )


# ---------------------------------------------------------------------------
# 마스터 대시보드: 고객사별 진행·유지보수 현황 (마스터 전용)
# ---------------------------------------------------------------------------
@router.get("/companies/overview")
async def companies_overview(
    ctx: AdminContext = Depends(require_super_admin),
    db: AsyncSession = Depends(get_db),
):
    """고객사마다 대상·완료·진행 중·중단·미시작·완료율, 리포트 미생성(유지보수 필요), 최근 활동, 준비 상태."""
    from sqlalchemy import func

    from diag_project.models.admin_user import AdminUser
    from diag_project.models.diagnosis_report import DiagnosisReport
    from diag_project.models.diagnosis_session import DiagnosisSession
    from diag_project.services.org_units import participant_statuses

    companies = (await db.execute(select(Company).order_by(Company.name))).scalars().all()
    hr_counts = dict((await db.execute(
        select(AdminUser.company_id, func.count(AdminUser.id)).where(AdminUser.company_id.is_not(None))
        .group_by(AdminUser.company_id)
    )).all())
    unit_counts = dict((await db.execute(
        select(OrgUnit.company_id, func.count(OrgUnit.id)).where(OrgUnit.is_active.is_(True)).group_by(OrgUnit.company_id)
    )).all())
    rows = []
    for c in companies:
        ps = (await db.execute(select(Participant.id, Participant.access_code).where(Participant.company_id == c.id))).all()
        pids = [p[0] for p in ps]
        st = await participant_statuses(db, pids)
        counts = {k: 0 for k in ("completed", "in_progress", "abandoned", "not_started")}
        for v in st.values():
            counts[v] += 1
        missing = 0
        last = None
        if pids:
            done_sessions = set((await db.execute(
                select(DiagnosisSession.user_id).where(DiagnosisSession.user_id.in_(pids), DiagnosisSession.status == "completed")
            )).scalars().all())
            with_report = set((await db.execute(
                select(DiagnosisReport.user_id).where(DiagnosisReport.user_id.in_(pids))
            )).scalars().all())
            # 완주했는데 리포트가 없는 사람 = 분석 실패 등 유지보수가 필요한 사람
            missing = len({str(u) for u in done_sessions} - {str(u) for u in with_report})
            last = (await db.execute(
                select(func.max(DiagnosisSession.updated_at)).where(DiagnosisSession.user_id.in_(pids))
            )).scalar()
        reg = len(pids)
        rows.append({
            "id": str(c.id), "name": c.name, "code": c.code, "org_prefix": c.org_prefix, "is_active": c.is_active,
            "diagnosis_start": c.diagnosis_start, "diagnosis_end": c.diagnosis_end,
            "registered": reg, "roster": sum(1 for p in ps if p[1]), **counts,
            "completion_rate": round(counts["completed"] / reg * 100) if reg else 0,
            "report_missing": missing, "last_activity": last,
            "hr_accounts": int(hr_counts.get(c.id, 0)), "org_unit_count": int(unit_counts.get(c.id, 0)),
        })
    return {"companies": rows}
